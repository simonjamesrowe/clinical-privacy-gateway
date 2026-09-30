"""Rebuild synthetic import fixtures locally; requires python-docx and reportlab.
No network, application database, environment credentials or real records are used.
"""
from pathlib import Path
from xml.sax.saxutils import escape
from datetime import datetime, timezone
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import reportlab
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4

ROOT = Path(__file__).resolve().parent
DISCLAIMER = 'SYNTHETIC TEST DATA - fictional people, places and clinical scenarios. Not for clinical use.'
# Invented combinations; .invalid domains and reserved fictional mobile numbers.
CASES = [
 ('sleep-review', 'Alex Morgan', 'sleep follow-up', 'Reported taking longer to fall asleep after changing work shifts. Described waking twice on most nights and feeling tired during the morning. No medication change was recorded.', 'Reported a more regular bedtime and fewer awakenings. Still felt tired after late shifts. The note records a plan to continue a sleep diary and discuss it at the next appointment.', 'partner Jamie Morgan', 'sleep diary and a four-week review'),
 ('anxiety-review', 'Priya Vale', 'anxiety follow-up', 'Described increased worry before travelling to work and difficulty concentrating during meetings. The source records practising a breathing exercise; it does not record a diagnosis or a risk assessment.', 'Reported using the exercise on three occasions and completing two journeys independently. Worry remained present before a busy meeting. No new physical symptoms were documented.', 'sibling Nila Vale', 'review of agreed coping strategies in three weeks'),
 ('pain-review', 'Daniel Wren', 'persistent pain review', 'Reported intermittent lower-back discomfort after prolonged sitting. Walking remained possible, with rests. The note does not include examination findings, imaging or a treatment recommendation.', 'Reported that shorter sitting periods helped during the week. Symptoms continued after a long car journey. A medicines list was not provided and must not be inferred.', 'support worker Rowan Ellis', 'functional goals and a planned review with the referring team'),
 ('mood-review', 'Samira Quill', 'mood and activity review', 'Described reduced motivation and withdrawing from a weekly craft group. Appetite was described as unchanged. There is no documented formal mental-state examination or safety assessment.', 'Reported attending one group session and enjoying a short walk with a friend. Motivation varied through the week. The source asks that uncertainty is retained rather than converted into diagnostic language.', 'friend Morgan Bell', 'activity diary and discussion at the next booked visit'),
 ('rehab-review', 'Theo Fenwick', 'rehabilitation progress', 'Recorded difficulty carrying shopping and managing stairs following a fictional period of reduced mobility. The baseline task time was 12 minutes with two rests; no clinical test score was recorded.', 'Recorded completion of the same task in 10 minutes with one rest. The patient described increased confidence but still requested help with heavier bags. These are reported observations, not evidence of recovery.', 'carer Ellis Fenwick', 'review of task pacing and practical support'),
 ('medicines-review', 'Zoë Marlow', 'medicines discussion', 'The synthetic note lists a previously prescribed 10 mg dose, with the drug name deliberately omitted. The patient reported taking it as listed and asked a question about timing. No prescribing decision was made.', 'The patient brought the same question to the second visit. A request for clarification was recorded for the prescribing clinician. The writer must not invent a drug name, titration, or clinical recommendation.', 'partner René Marlow', 'clarification from the prescribing clinician'),
 ('family-meeting', 'Casey Alder', 'family support meeting', 'A joint appointment included a relative also named Casey Alder. The patient discussed difficulty arranging transport; the relative described being available on Tuesdays. Both names intentionally match for placeholder testing.', 'The patient reported one successful journey. The relative could not attend and their availability was not updated. Keep patient and relative statements separate even when visible names are the same.', 'relative Casey Alder', 'transport options and consent to future family involvement'),
 ('work-review', 'Jules Hartwell', 'work participation review', 'Described fatigue during a fictional office role at Cedar Lantern Studio. The manager was named Taylor Moss. The note records consent to discuss functional needs but does not record consent to disclose a diagnosis.', 'Reported trying shorter meetings and scheduled breaks. The patient wanted an employer-facing summary containing functional observations only. Hours of work and legal eligibility were not recorded.', 'manager Taylor Moss', 'a functional summary with no unsupported certification'),
 ('referral-notes', 'Leila Brook', 'referral preparation', 'Requested a specialist opinion about a persistent concern documented only as difficulty with daily activities. Symptom onset, examination findings and referral urgency were not recorded.', 'Reported that the concern remained and asked how referral information would be shared. The note lists prior strategies as a diary and a routine review. There is no documented specialist diagnosis.', 'GP Dr Avery Reed', 'a referral request clearly identifying missing information'),
 ('discharge-review', 'Owen Thistle', 'end-of-episode review', 'At the start of the fictional episode, the patient described difficulty establishing a weekly routine. Agreed goals were attending a community activity and keeping a brief progress diary.', 'At the final visit, the patient reported attending two activities and keeping the diary for five days. Some goals remained unfinished. The record notes an agreed end to this episode, without asserting that all symptoms resolved.', 'community worker Erin Thistle', 'a balanced summary of progress and unresolved goals'),
]
TEMPLATES = [
 ('gp-update', 'GP update letter', 'Write a concise update to the GP about the selected reviewed notes.', ['Reason for contact', 'Relevant history from the notes', 'Progress', 'Agreed next steps']),
 ('referral-letter', 'Referral letter', 'Draft a request for an opinion; do not invent a referral urgency or specialist diagnosis.', ['Reason for referral', 'Documented presentation', 'Steps already taken', 'Questions for the receiving clinician']),
 ('progress-report', 'Progress report', 'Compare earlier and later observations, preserving dates and uncertainty.', ['Starting point', 'Changes over time', 'Ongoing difficulties', 'Plan recorded in the notes']),
 ('patient-summary', 'Patient summary', 'Use accessible British English for a patient-facing summary, without adding advice.', ['What we discussed', 'Progress you reported', 'What was agreed', 'Questions still to discuss']),
 ('multidisciplinary', 'Team handover', 'Summarise only documented information for a multidisciplinary team.', ['Current situation', 'Relevant background', 'Reported observations', 'Outstanding actions and recorded owners']),
 ('functional-update', 'Functional update', 'Describe functional activities and support needs; omit unnecessary diagnostic labels.', ['Purpose', 'Activities described', 'Support tried', 'Remaining needs']),
 ('medicines-query', 'Medicines clarification', 'Draft a clarification request to the prescribing clinician; never supply missing prescribing details.', ['Recorded medicine information', 'Patient question', 'Documented observations', 'Clarification requested']),
 ('family-meeting', 'Family meeting summary', 'Keep patient and relative statements separate, including people with repeated names.', ['Participants as recorded', 'Patient perspective', 'Relative perspective', 'Agreed actions and consent limits']),
 ('episode-summary', 'End-of-episode summary', 'Describe progress and unresolved goals without claiming recovery or inventing follow-up.', ['Reason for involvement', 'Work documented', 'Progress against goals', 'Outstanding matters']),
 ('chronology', 'Clinical chronology', 'Produce a dated chronology from the supplied notes, keeping conflicting accounts visible.', ['Dated events', 'Changes recorded', 'Conflicting or missing information', 'Latest documented plan']),
]

def sections(index, case):
    slug, name, topic, first, second, contact, plan = case
    return [
      ('Patient details', f'Fictional patient: {name}. Test reference: SYN-{index:04d}. Date of birth: {index:02d}/04/1980. Address: {index} Imaginary Mews, Exampleford, XX1 1XX. Email: case{index:02d}@example.invalid. Telephone: 07700 900{index:03d}.'),
      ('Visit 1 - 03 September 2026', f'{name} attended the fictional Exampleford Practice for a {topic}. {first} The record names {contact}. Identifying details are included deliberately so that local detection and review can be exercised.'),
      ('Visit 2 - 24 September 2026', f'{name} attended a second appointment with fictional clinician Dr Avery Reed. {second} The documented next step is {plan}. The phrase [CLIENT] appears here as literal source text for substitution testing.'),
      ('Information boundaries', 'This scenario supplies no real diagnosis, risk conclusion, examination result or treatment recommendation. Missing information must remain missing. The two visits are separate observations and should not be silently combined into a single event.'),
    ]

def word(path, title, parts):
    doc = Document(); section = doc.sections[0]
    # standard_business_brief, memo_masthead; A4 override for UK clinical test files.
    section.page_width = Inches(8.2677); section.page_height = Inches(11.6929)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.492)
    for name, size, before, after, colour in [('Normal',11,0,6,'122426'),('Title',22,0,8,'122426'),('Subtitle',11,0,6,'496062'),('Heading 1',16,16,8,'2E74B5'),('Heading 2',13,12,6,'2E74B5'),('Heading 3',12,8,4,'1F4D78')]:
        style=doc.styles[name]; style.font.name='Calibri'; style.font.size=Pt(size); style.font.color.rgb=RGBColor.from_string(colour)
        fmt=style.paragraph_format; fmt.space_before=Pt(before);fmt.space_after=Pt(after);fmt.line_spacing=1.10
        if name.startswith('Heading'): fmt.keep_with_next=True
    doc.core_properties.author='Synthetic test fixture generator'; doc.core_properties.last_modified_by='Synthetic test fixture generator'
    doc.core_properties.title=title; doc.core_properties.created=datetime(2026,9,29,tzinfo=timezone.utc); doc.core_properties.modified=datetime(2026,9,29,tzinfo=timezone.utc)
    section.header.paragraphs[0].text='CLINICIAN’S VEIL / SYNTHETIC TEST FIXTURE'
    footer=section.footer.paragraphs[0]; footer.text='Fictional content only | Page '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.add_paragraph(title,'Title');doc.add_paragraph(DISCLAIMER,'Subtitle')
    for heading, content in parts:
        doc.add_paragraph(heading,'Heading 2');doc.add_paragraph(content)
    doc.save(path)

def pdf(path, title, parts):
    # Embed Unicode maps instead of relying on OS-specific standard-font decoding.
    fonts = Path(reportlab.__file__).parent / 'fonts'
    pdfmetrics.registerFont(TTFont('FixtureSans', str(fonts / 'Vera.ttf')))
    pdfmetrics.registerFont(TTFont('FixtureSansBold', str(fonts / 'VeraBd.ttf')))
    body=ParagraphStyle('Body',fontName='FixtureSans',fontSize=10.5,leading=14,spaceAfter=8)
    heading=ParagraphStyle('Heading',parent=body,fontName='FixtureSansBold',fontSize=12,leading=15,spaceBefore=12,spaceAfter=6,keepWithNext=True)
    title_style=ParagraphStyle('Title',parent=heading,fontSize=21,leading=25,spaceBefore=0)
    story=[Paragraph(escape(title),title_style),Paragraph(escape(DISCLAIMER),body),Spacer(1,8)]
    for label,content in parts: story.extend([Paragraph(escape(label),heading),Paragraph(escape(content),body)])
    def footer(canvas, doc):
        canvas.setFont('FixtureSans',8);canvas.drawString(54,30,'SYNTHETIC TEST FIXTURE - fictional content only');canvas.drawRightString(A4[0]-54,30,str(doc.page))
    SimpleDocTemplate(str(path),pagesize=A4,leftMargin=54,rightMargin=54,topMargin=45,bottomMargin=48,title=title,author='Synthetic test fixture generator').build(story,onFirstPage=footer,onLaterPages=footer)

def main():
    (ROOT / "FONT-LICENSE.txt").write_text((Path(reportlab.__file__).parent / "fonts/bitstream-vera-license.txt").read_text())
    for folder in ['word','pdf','notes','templates']: (ROOT/folder).mkdir(exist_ok=True)
    manifest=[]
    for index,case in enumerate(CASES,1):
        stem=f'{index:02d}-{case[0]}';title=f'Synthetic {case[2]}';parts=sections(index,case)
        (ROOT/'notes'/f'{stem}.txt').write_text(DISCLAIMER+'\n\n'+title+'\n\n'+'\n\n'.join(f'{heading}\n{content}' for heading,content in parts)+'\n')
        word(ROOT/'word'/f'{stem}.docx',title,parts);pdf(ROOT/'pdf'/f'{stem}.pdf',title,parts)
        manifest.append({'id':f'SYN-{index:04d}','name':case[1],'topic':case[2],'stem':stem,'synthetic':True})
    for index,(slug,title,purpose,headings) in enumerate(TEMPLATES,1):
        text=f'# {title}\n\n## Purpose\n\n{purpose}\n\n## Source rules\n\n- Use **only** the supplied reviewed notes.\n- Preserve every request placeholder exactly; never guess an identity.\n- Do not invent facts, diagnoses, risk findings, medication details or actions.\n- State that information is not recorded when it is missing.\n- Distinguish reported information from recorded observations.\n\n## Document structure\n\n'+ '\n'.join(f'- **{heading}**' for heading in headings)+'\n\n## Style\n\nUse professional British English, short paragraphs and clear headings. Do not add a signature, clinician details or contact information; these are handled locally.\n'
        (ROOT/'templates'/f'{index:02d}-{slug}.md').write_text(text)
    (ROOT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print('Created 10 Word documents, 10 PDFs, 10 text notes and 10 Markdown prompt templates.')
if __name__ == '__main__': main()
